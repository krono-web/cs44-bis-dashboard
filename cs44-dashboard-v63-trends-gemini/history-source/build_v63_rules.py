"""Rebuild the 48-rule v6.3 configuration from the two client workbooks.

The three classifier additions documented by the adjudication/results are supplied
explicitly below. The output must reproduce the original 2026 dashboard exactly
before it is used for another year.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


CATEGORY_CODES = {
    "Work-Integrated and Applied Learning": "work_integrated_applied",
    "Simulation": "simulation",
    "Case-Based Learning": "case_based",
    "Project- and Problem-Based Learning": "project_problem_based",
    "Community Learning": "community_learning",
    "Entrepreneurial Learning": "entrepreneurial_learning",
    "Technology-Mediated Learning": "technology_mediated_learning",
    "Hybrid Learning": "hybrid_learning",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--changes", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    matrix = load_workbook(args.matrix, read_only=True, data_only=True)
    changes = load_workbook(args.changes, read_only=True, data_only=True)
    rules: dict[str, dict] = {}
    for number, row in enumerate(list(matrix["Scoring Matrix"].values)[4:], 1):
        category, label, phrases, weight, treatment, *_ = row
        rules[f"R{number:03}"] = {
            "id": f"R{number:03}", "category": CATEGORY_CODES[category],
            "label": label, "pattern": "|".join(x.strip() for x in phrases.split(";")),
            "weight": float(weight), "review_required": treatment == "Review only",
        }

    for row in list(changes["Before and after"].values)[1:]:
        rule_id, after = row[1], row[4]
        lines = after.splitlines()
        label = re.sub(r"\s+\(was:.*", "", lines[0]).strip()
        category = re.sub(r"\s+\(was:.*", "", lines[1]).strip()
        weight = float(re.search(r"weight ([\d.]+)", lines[2]).group(1))
        phrases = [
            re.sub(r"\s+\(added by us\)", "", value).lstrip("+ ").strip()
            for value in after.split("\n\n", 1)[1].splitlines() if value.strip()
        ]
        rules[rule_id] = {
            "id": rule_id, "category": CATEGORY_CODES[category], "label": label,
            "pattern": "|".join(phrases), "weight": weight,
            "review_required": lines[3] == "review-only",
        }

    # The change explanations define these two cross-cutting review rules.
    rules["R044"] = {
        "id": "R044", "category": "work_integrated_applied",
        "label": "generic professional skills",
        "pattern": "|".join([
            "experiential knowledge", "applied knowledge", "in practice", "scenario",
            "hands-on activities", "hands-on", "practical skill", "teamwork",
            "communication", "stakeholder", "leadership", "self-development",
            "development", "critical thinking", "problem-solving", "discovery",
            "real-world", "real world", "immersion",
        ]),
        "weight": 0.5, "review_required": True,
    }
    rules["R045"] = {
        "id": "R045", "category": "project_problem_based",
        "label": "generic strategy term", "pattern": "strategy",
        "weight": 2.0, "review_required": True,
    }
    # These three rules are evidenced by the adjudication sheet and uniquely
    # recoverable from the original data.js matched phrases and item scores.
    rules["R046"] = {
        "id": "R046", "category": "simulation", "label": "simulation activity",
        "pattern": "simulation|simulate|simulated|simulating",
        "weight": 4.0, "review_required": False,
    }
    rules["R047"] = {
        "id": "R047", "category": "technology_mediated_learning",
        "label": "generative ai mention", "pattern": "generative ai|genai",
        "weight": 2.0, "review_required": True,
    }
    rules["R048"] = {
        "id": "R048", "category": "technology_mediated_learning",
        "label": "ai-enabled learning activity", "pattern": "ai-enabled|ai enabled",
        "weight": 4.0, "review_required": False,
    }

    category_definitions = {
        "work_integrated_applied": "An educational approach that explicitly merges theory with real-world practice and embeds authentic industry, workplace, or community-relevant work and tasks into a unit of study.",
        "simulation": "An instructional method that recreates or mimics real-world scenarios, environments, or processes so learners can practise skills and make decisions safely.",
        "case_based": "An active teaching method in which students use real-world scenarios and case studies to build critical-thinking and problem-solving skills.",
        "project_problem_based": "Learning focused on real-world problems and challenges using problem-solving, decision-making, investigative, and self-directed learning skills.",
        "community_learning": "Learning undertaken together with or for communities, including service-learning, field studies, and professional communities of practice.",
        "entrepreneurial_learning": "Learning that develops the knowledge, skills, and mindset to recognise opportunities, turn ideas into action, and manage new ventures under uncertainty.",
        "technology_mediated_learning": "An instructional environment in which digital devices, networks, or specialised technologies mediate access to content and interaction during learning.",
        "hybrid_learning": "An education model that mixes face-to-face classroom teaching with online digital instruction.",
    }
    categories = []
    for name, code in CATEGORY_CODES.items():
        category_rules = []
        for rule in rules.values():
            if rule["category"] != code:
                continue
            category_rules.append({k: v for k, v in rule.items() if k != "category"})
        categories.append({"id": code, "name": name, "definition": category_definitions[code], "rules": category_rules})

    payload = {
        "version": "2026-10-07-v6.3-client-review",
        "title": "CS-44 Pedagogical Innovation Taxonomy v6.3",
        "counting_unit": "One distinct outline item evaluated once per category.",
        "decision_thresholds": {"positive": 3.0, "high_confidence": 5.0, "status": "provisional_pending_validation"},
        "policies": {"multi_label": True, "categories_mutually_exclusive": False, "deduplication": "item-category", "rule_alternatives_count_once": True, "review_only_cannot_classify": True},
        "provenance": {
            "matrix_sha256": sha256(args.matrix), "changes_sha256": sha256(args.changes),
            "reconstruction_note": "43-row v5 matrix updated by the 19-row client change log; five added rules recovered from its explanations, the adjudication workbook, and exact original data.js record-level scores. Must pass verify_v63.py before use.",
        },
        "categories": categories,
    }
    if sum(len(c["rules"]) for c in categories) != 48:
        raise ValueError("Expected exactly 48 rules")
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
